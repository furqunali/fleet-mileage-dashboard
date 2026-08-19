#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_console.py  --  create the single HR master workbook.

Generates  Reimbursement_Workflow/Gas_Reimbursement_Console.xlsx  -- the ONE
file HR ever opens. Tabs:

  Dashboard      -- month selector + KPIs + all-month history (engine-filled)
  Monthly Input  -- HR pastes every month's trips here (all history kept)
  Employees      -- roster + rates
  Locations      -- verified HOME / WORK / OTHER locations + coordinates
  Settings       -- policy knobs (LOCKED so HR can't change by accident)
  Monthly Report -- per-employee summary for the selected month (engine)
  Trip Detail    -- every leg for the selected month (engine)
  Exceptions     -- everything needing review for the selected month (engine)
  Audit Trail    -- run metadata (engine)

The master Directory/Matrix stays a SEPARATE, protected, read-only file; the
Python engine (process_console.py) is the backend. Matrix miles are
authoritative; locations not on the Matrix are flagged for review and never
auto-paid.

Seeds the roster/locations and migrates the existing May-2026 sample into the
Monthly Input tab so the workbook is immediately usable and QA-comparable.

SAFETY: refuses to overwrite an existing console unless --force.
Usage:  python build_console.py [--force]
"""

import argparse
import os
import sys

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
if SCRIPT_DIR not in sys.path:
    sys.path.insert(0, SCRIPT_DIR)
import console_layout as L
import process_month as pm   # reuse the tested input parser for migration

WORKFLOW_DIR = os.path.dirname(SCRIPT_DIR)
OUT_PATH = os.path.join(WORKFLOW_DIR, L.CONSOLE_FILENAME)
SEED_MONTH = "2026-05"
SEED_RAW = os.path.join(WORKFLOW_DIR, "inbox", SEED_MONTH,
                        "2026-05 Gas Reimbursement - Employee A.xlsx")

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

# ---- styling ----
BRAND = "1F4E78"
HDR_FILL = PatternFill("solid", fgColor=BRAND)
HDR_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(size=16, bold=True, color=BRAND)
SUB_FONT = Font(size=11, italic=True, color="595959")
BAND = PatternFill("solid", fgColor="FFC000")
INPUT_FILL = PatternFill("solid", fgColor="FFF7E6")   # pale amber = "type here"
LOCK_FILL = PatternFill("solid", fgColor="EDEDED")    # grey = locked/engine
KPI_FILL = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def hdr(ws, headers, row=1, comments=None):
    for ci, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=ci, value=h)
        c.fill = HDR_FILL; c.font = HDR_FONT
        c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        c.border = BORDER
        if comments and h in comments:
            c.comment = Comment(comments[h], "HR Console")
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = f"A{row + 1}"


def dropdown(ws, col, options=None, formula=None, first=2, last=600):
    f1 = ('"' + ",".join(options) + '"') if options else formula
    dv = DataValidation(type="list", formula1=f1, allow_blank=True)
    dv.prompt = "Pick from the list"
    ws.add_data_validation(dv)
    dv.add(f"{col}{first}:{col}{last}")


def widths(ws, mapping):
    for k, v in mapping.items():
        ws.column_dimensions[k].width = v


def fill_input_region(ws, ncols, first=2, last=None, buffer=30):
    """Apply the pale 'type here' fill + borders. If `last` is None, fill the
    used rows plus a small buffer of blank rows (avoids inflating max_row to
    hundreds, which would push ws.append() far down the sheet)."""
    if last is None:
        last = ws.max_row + buffer
    for r in range(first, last + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            cell.fill = INPUT_FILL


# ------------------------------------------------------------------ builders
def build_dashboard(ws):
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Fleet Gas Reimbursement — HR Console"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = ("One workbook. Paste trips on Monthly Input, keep Employees & "
                "Locations current, pick a month below, then run.")
    ws["A2"].font = SUB_FONT

    ws["B3"] = "Selected Month:"; ws["B3"].font = Font(bold=True)
    sel = ws[L.DASH_SELECTED_MONTH_CELL]
    sel.value = SEED_MONTH
    sel.fill = INPUT_FILL; sel.border = BORDER
    sel.font = Font(bold=True, color=BRAND)
    ws["B4"] = "Last run:"; ws["B4"].font = Font(bold=True)
    ws[L.DASH_LAST_RUN_CELL] = "(not run yet)"

    # month selector dropdown (engine refreshes the list each run)
    dv = DataValidation(type="list", formula1='"' + SEED_MONTH + '"', allow_blank=True)
    dv.prompt = "Pick a month that exists on Monthly Input"
    ws.add_data_validation(dv)
    dv.add(f"{L.DASH_SELECTED_MONTH_CELL}")

    # KPI block (labels now; engine fills values)
    ws.cell(row=L.DASH_KPI_START_ROW - 1, column=2, value="THIS MONTH").font = Font(bold=True, color=BRAND)
    for i, key in enumerate(L.DASH_KPI_KEYS):
        r = L.DASH_KPI_START_ROW + i
        lc = ws.cell(row=r, column=2, value=key); lc.font = Font(bold=True)
        lc.border = BORDER
        vc = ws.cell(row=r, column=3); vc.fill = KPI_FILL; vc.border = BORDER

    # history table header (engine fills rows)
    ws.cell(row=L.DASH_HISTORY_HEADER_ROW - 1, column=2,
            value="ALL MONTHS (history)").font = Font(bold=True, color=BRAND)
    for ci, h in enumerate(L.DASH_HISTORY_HEADERS, start=2):
        c = ws.cell(row=L.DASH_HISTORY_HEADER_ROW, column=ci, value=h)
        c.fill = HDR_FILL; c.font = HDR_FONT; c.border = BORDER

    # legend
    lr = L.DASH_HISTORY_HEADER_ROW + L.DASH_HISTORY_MAX_ROWS + 2
    ws.cell(row=lr, column=2, value="How to run:").font = Font(bold=True, color=BRAND)
    for j, txt in enumerate([
        "1) Fill Monthly Input / Employees / Locations.",
        "2) Save and CLOSE this workbook.",
        "3) Double-click Run.bat.",
        "4) Re-open — Monthly Report, Trip Detail, Exceptions and this Dashboard are updated.",
        "Green = approved · Amber = needs review · Red = exception.",
        "Matrix miles are authoritative. Locations not on the Matrix are never auto-paid.",
    ]):
        ws.cell(row=lr + 1 + j, column=2, value=txt).font = SUB_FONT
    widths(ws, {"A": 4, "B": 26, "C": 22, "D": 12, "E": 12, "F": 12, "G": 12, "H": 12})


def build_input(ws):
    hdr(ws, L.INPUT_HEADERS, comments={
        "Month (YYYY-MM)": "e.g. 2026-05. All months live here; the Dashboard selector picks which one to report.",
        "Employee ID": "Must exist on the Employees tab.",
        "Claimed Miles": "Optional. If blank, computed as Odometer End - Odometer Start.",
    })
    # NOTE: fill applied AFTER migration in build(), so appended rows land at row 2.
    dropdown(ws, "B", formula=f"{L.SH_EMPLOYEES}!$A$2:$A$200")
    widths(ws, {"A": 15, "B": 12, "C": 12, "D": 30, "E": 30, "F": 22,
                "G": 15, "H": 15, "I": 13})


def build_employees(ws):
    hdr(ws, L.EMPLOYEE_HEADERS, comments={
        "Employee ID": "Unique ID used across the workbook.",
        "Default Rate": "Per-mile rate; falls back to Settings if blank.",
        "Active": "YES = process. NO = keep but skip.",
    })
    ws.append(["1001", "JANE DOE", 0.27, "YES",
               "Home unconfirmed — see Locations. Home legs stay $0 until HR approval."])
    fill_input_region(ws, len(L.EMPLOYEE_HEADERS))
    dropdown(ws, "D", options=["YES", "NO"])
    widths(ws, {"A": 14, "B": 26, "C": 13, "D": 10, "E": 55})


def build_locations(ws):
    hdr(ws, L.LOCATION_HEADERS, comments={
        "Employee ID": "Whose location. Use ALL for a shared location.",
        "Type": "HOME (always $0/review), WORK, or OTHER.",
        "Matrix Site Code": "Link to a Directory site code (e.g. 0083) so official Matrix miles are used.",
        "HR Approved": "YES after confirming. Home legs still show $0 pending approval.",
    })
    ws.append(["LOC001", "1001", "HOME",
               "demo apartment;demo apartmen;aparment;aparmen;aparment houston;"
               "aparmen houston;home;houston home;home houston;aparmen stores",
               "", "", "", "", "NO",
               "Confirm which Houston 'Demo Apartments' then fill Address + Lat/Lon. Reference only; $0 until approval."])
    ws.append(["LOC002", "ALL", "WORK",
               "123 demo blvd;123 demo blvd;123 demo blvd",
               "123 Demo Blvd, Democity, TX 70001", "", "", "0075", "YES",
               "Verified alias -> Demo Store North (0075). Matrix miles authoritative."])
    ws.append(["LOC003", "ALL", "WORK", "west store",
               "456 Sample Rd, Demotown, TX 70002", "", "", "0083", "YES",
               "Verified alias -> Demo Store West (0083). Matrix miles authoritative."])
    fill_input_region(ws, len(L.LOCATION_HEADERS))
    dropdown(ws, "C", options=["HOME", "WORK", "OTHER"])
    dropdown(ws, "I", options=["YES", "NO"])
    widths(ws, {"A": 11, "B": 11, "C": 8, "D": 42, "E": 38, "F": 11, "G": 11,
                "H": 15, "I": 12, "J": 48})


def build_settings(ws):
    hdr(ws, ["Setting", "Value", "Description"])
    for key, label, default, desc in L.SETTINGS_ROWS:
        ws.append([key, default, desc])
        ws.cell(row=ws.max_row, column=1).comment = Comment(label, "HR Console")
    for r in range(2, ws.max_row + 1):
        for c in range(1, 4):
            ws.cell(row=r, column=c).border = BORDER
            ws.cell(row=r, column=c).fill = LOCK_FILL
    widths(ws, {"A": 26, "B": 34, "C": 60})
    # LOCK the sheet so HR cannot change policy by accident.
    ws.protection.sheet = True
    ws.protection.enable()
    ws["A1"].comment = Comment("This tab is locked. To change policy, unprotect the sheet first.",
                               "HR Console")


def build_output_placeholder(ws, headers, note):
    hdr(ws, headers)
    ws.cell(row=2, column=1, value=note).font = SUB_FONT
    for c in range(1, len(headers) + 1):
        ws.cell(row=1, column=c)


def migrate_seed(ws_input):
    """Populate Monthly Input from the existing May raw file (exact same parse
    the engine uses), so the console is immediately usable + QA-comparable."""
    if not os.path.exists(SEED_RAW):
        print(f"[warn] seed raw file not found, Monthly Input left empty: {SEED_RAW}")
        return 0
    policy = {"master_workbook": "x"}  # parse_input ignores master; only needs sheet parse
    ef = pm.parse_input(SEED_RAW, policy)
    n = 0
    for t in ef.trips:
        d = t.date
        try:
            dstr = d.strftime("%Y-%m-%d")
        except AttributeError:
            dstr = str(d).split(" ")[0]
        ws_input.append([SEED_MONTH, ef.emp_id or "1001", dstr, t.origin_raw, t.dest_raw,
                         t.notes, t.odo_start, t.odo_end, t.claimed_miles])
        n += 1
    fill_input_region(ws_input, len(L.INPUT_HEADERS))   # used rows + blank buffer
    return n


def build():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    sheets = {name: wb.create_sheet(name) for name in L.SHEET_ORDER}

    build_dashboard(sheets[L.SH_DASHBOARD])
    build_input(sheets[L.SH_INPUT])
    build_employees(sheets[L.SH_EMPLOYEES])
    build_locations(sheets[L.SH_LOCATIONS])
    build_settings(sheets[L.SH_SETTINGS])
    build_output_placeholder(sheets[L.SH_REPORT], L.REPORT_HEADERS, "(Run to populate.)")
    build_output_placeholder(sheets[L.SH_TRIPS], L.TRIP_HEADERS, "(Run to populate.)")
    build_output_placeholder(sheets[L.SH_EXCEPTIONS], L.TRIP_HEADERS, "(Run to populate.)")
    build_output_placeholder(sheets[L.SH_AUDIT], L.AUDIT_HEADERS, "(Run to populate.)")

    n = migrate_seed(sheets[L.SH_INPUT])
    wb.active = wb.sheetnames.index(L.SH_DASHBOARD)
    wb.save(OUT_PATH)
    print(f"[ok] Wrote {OUT_PATH}")
    print(f"[ok] Migrated {n} May-2026 rows into Monthly Input.")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--force", action="store_true",
                    help="Overwrite an existing console with a fresh template (destroys HR data).")
    args = ap.parse_args()
    if os.path.exists(OUT_PATH) and not args.force:
        sys.exit(f"[skip] {OUT_PATH} already exists. Use --force to overwrite (destroys HR data).")
    build()


if __name__ == "__main__":
    main()
