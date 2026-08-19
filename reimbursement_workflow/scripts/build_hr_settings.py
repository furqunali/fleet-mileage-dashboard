#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_hr_settings.py  --  create the HR control-panel workbook.

Generates  config/HR_Settings.xlsx  -- the single friendly place where HR
maintains, from their own computer, the data the monthly workflow needs:

  * Employees      -- roster (id, name, rate, active)
  * Locations      -- verified home / work / other addresses + coordinates,
                      fully expandable: add employees, add many locations each,
                      edit addresses, update coordinates over time.
  * Instructions   -- how to use it.

The monthly engine (process_month.py) reads this workbook every run, so any row
HR adds or edits is picked up automatically for the next month -- no code change.
The master Directory/Matrix and the audit trail are never touched by HR edits.

SAFETY: this script will NOT overwrite an existing HR_Settings.xlsx (that would
destroy HR's data). Pass --force only to regenerate a blank template.

Usage:
    python build_hr_settings.py            # create if missing
    python build_hr_settings.py --force    # overwrite with a fresh template
"""

import argparse
import os
import sys

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
CONFIG_DIR = os.path.join(os.path.dirname(SCRIPT_DIR), "config")
OUT_PATH = os.path.join(CONFIG_DIR, "HR_Settings.xlsx")

# --- styling ---
BRAND = "1F4E78"
HDR_FILL = PatternFill("solid", fgColor=BRAND)
HDR_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(size=15, bold=True, color=BRAND)
NOTE_FONT = Font(italic=True, color="595959")
BAND = PatternFill("solid", fgColor="FFC000")
INPUT_FILL = PatternFill("solid", fgColor="FFF7E6")   # pale = "type here"
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(ws, headers, row=1, comments=None):
    for ci, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=ci, value=h)
        c.fill = HDR_FILL
        c.font = HDR_FONT
        c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        c.border = BORDER
        if comments and h in comments:
            c.comment = Comment(comments[h], "HR Settings")
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def add_dropdown(ws, col_letter, options, first_row=2, last_row=400):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"',
                        allow_blank=True, showDropDown=False)
    dv.error = "Pick a value from the list."
    dv.prompt = "Choose: " + ", ".join(options)
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}{first_row}:{col_letter}{last_row}")


def widths(ws, mapping):
    for letter, w in mapping.items():
        ws.column_dimensions[letter].width = w


def build():
    wb = openpyxl.Workbook()

    # ================= Instructions =================
    ws = wb.active
    ws.title = "Instructions"
    ws["A1"] = "Fleet Gas-Reimbursement — HR Settings"
    ws["A1"].font = TITLE_FONT
    for cc in range(1, 3):
        ws.cell(row=2, column=cc).fill = BAND
    lines = [
        "",
        "This workbook is the control panel for the monthly gas-reimbursement workflow.",
        "Edit it from your own computer. The workflow reads it every time it runs, so",
        "anything you add or change here is used automatically for the next month —",
        "no programmer needed.",
        "",
        "TABS:",
        "  • Employees — one row per employee. Add new staff here.",
        "  • Locations — verified home / work addresses with coordinates. Add as many",
        "    rows as you like, for any employee, now or in the future.",
        "",
        "HOW TO ADD OR UPDATE A HOME ADDRESS:",
        "  1. Go to the Locations tab.",
        "  2. Find (or add) a row for the employee (use their Employee ID).",
        "  3. Set Type = HOME.",
        "  4. In 'Aliases', list how the address appears in the raw log, separated by",
        "     semicolons  (e.g.  demo apartment; demo apartmen; home houston).",
        "  5. Enter the Verified Address.",
        "  6. Enter Latitude / Longitude from Google Maps (right-click the pin in",
        "     Google Maps → the first line copies the coordinates). Google Maps is a",
        "     REFERENCE ONLY — company Matrix miles remain the official distance.",
        "  7. Leave 'HR Approved' = NO until you have confirmed it with the employee.",
        "",
        "IMPORTANT POLICY (unchanged):",
        "  • Home / apartment legs are ALWAYS shown as HOME_ADDRESS_REVIEW with $0",
        "    reimbursement until HR approves them — even after the address is filled in.",
        "  • Multi-store lines are ALWAYS flagged for manual itemisation ($0, no auto",
        "    distance). Route order is never guessed.",
        "  • The official site Directory and distance Matrix are never changed here.",
        "",
        "COORDINATES let the report show a straight-line reference distance for a home",
        "or off-Matrix leg to help you review — it is a reference figure, not payment.",
        "",
        "TIP: To link a new WORK location to an existing company site so official Matrix",
        "miles are used, put that site's code (e.g. 0083) in 'Matrix Site Code'.",
    ]
    r = 3
    for ln in lines:
        c = ws.cell(row=r, column=1, value=ln)
        if ln.endswith(":") or ln.isupper():
            c.font = Font(bold=True, color=BRAND)
        else:
            c.font = NOTE_FONT if ln.startswith(("  ", "COORD", "This", "Edit", "any", "no ")) else Font()
        r += 1
    widths(ws, {"A": 95})

    # ================= Employees =================
    ws = wb.create_sheet("Employees")
    emp_hdr = ["Employee ID", "Employee Name", "Default Rate", "Active", "Notes"]
    emp_comments = {
        "Employee ID": "Must match the ID that appears inside the employee's monthly file.",
        "Default Rate": "Used only if the monthly file's own rate cell is blank (e.g. 0.27).",
        "Active": "YES = include in processing. NO = keep the record but skip.",
    }
    style_header(ws, emp_hdr, comments=emp_comments)
    ws.append(["1001", "JANE DOE", 0.27, "YES",
               "Home address unconfirmed — see Locations tab. Home legs stay $0 until HR approval."])
    add_dropdown(ws, "D", ["YES", "NO"])
    widths(ws, {"A": 14, "B": 26, "C": 13, "D": 10, "E": 60})
    for row in ws.iter_rows(min_row=2, max_row=200, min_col=1, max_col=5):
        for c in row:
            c.border = BORDER
            if c.column_letter in ("A", "B", "C", "D", "E"):
                c.fill = INPUT_FILL

    # ================= Locations =================
    ws = wb.create_sheet("Locations")
    loc_hdr = ["Location ID", "Employee ID", "Type", "Aliases (as in log; separate with ;)",
               "Verified Address", "Latitude", "Longitude", "Matrix Site Code",
               "HR Approved", "Notes"]
    loc_comments = {
        "Employee ID": "The employee this location belongs to. Use ALL for a shared location.",
        "Type": "HOME = residence (always $0/review). WORK = a work site. OTHER = anything else.",
        "Aliases (as in log; separate with ;)":
            "Every spelling that may appear in the raw log, separated by semicolons.",
        "Verified Address": "The confirmed street address (free text).",
        "Latitude": "From Google Maps (reference only). Decimal, e.g. 29.7016.",
        "Longitude": "From Google Maps (reference only). Decimal, e.g. -95.3543.",
        "Matrix Site Code": "Optional. Link to an existing company site code (e.g. 0083) to use "
                            "official Matrix miles for a WORK location.",
        "HR Approved": "YES only after you have confirmed the location. Home legs still show $0.",
    }
    style_header(ws, loc_hdr, comments=loc_comments)
    # Seed: Employee A's home aliases (address/coords blank, unconfirmed).
    ws.append(["LOC001", "1001", "HOME",
               "demo apartment;demo apartmen;aparment;aparmen;aparment houston;"
               "aparmen houston;home;houston home;home houston;aparmen stores",
               "", "", "", "", "NO",
               "Multiple Houston 'Demo Apartments' properties exist — confirm the exact one, then fill "
               "Address + Latitude/Longitude. Reference only; home legs remain $0 until approval."])
    # Example (grey) rows to show the pattern — HR can overwrite or delete.
    ws.append(["(example)", "1001", "WORK", "west store", "456 Sample Rd, Demotown, TX 70002",
               29.8179017, -95.739811, "0083", "YES",
               "Example: WORK location linked to site 0083 -> official Matrix miles used."])
    add_dropdown(ws, "C", ["HOME", "WORK", "OTHER"])
    add_dropdown(ws, "I", ["YES", "NO"])
    widths(ws, {"A": 12, "B": 12, "C": 9, "D": 40, "E": 40,
                "F": 12, "G": 12, "H": 14, "I": 12, "J": 50})
    for row in ws.iter_rows(min_row=2, max_row=400, min_col=1, max_col=10):
        for c in row:
            c.border = BORDER
            c.fill = INPUT_FILL
    # grey-out the example row
    for c in ws[3]:
        c.font = Font(italic=True, color="A6A6A6")

    wb.save(OUT_PATH)
    print(f"[ok] Wrote {OUT_PATH}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--force", action="store_true",
                    help="Overwrite an existing HR_Settings.xlsx with a blank template.")
    args = ap.parse_args()
    if os.path.exists(OUT_PATH) and not args.force:
        sys.exit(f"[skip] {OUT_PATH} already exists. Use --force to overwrite (destroys HR data).")
    os.makedirs(CONFIG_DIR, exist_ok=True)
    build()


if __name__ == "__main__":
    main()
