#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
process_console.py  --  backend engine for the single HR console workbook.

Reads the HR-maintained tabs (Monthly Input, Employees, Locations, Settings)
and the selected month from the Dashboard, matches every leg against the
authoritative (read-only) master Directory/Matrix, then writes the output tabs
(Monthly Report, Trip Detail, Exceptions, Audit Trail) and the Dashboard KPIs +
all-month history back into the SAME workbook.

Rules (unchanged, locked):
  * Matrix road miles are authoritative. Odometer is a cross-check only.
  * HOME/apartment legs -> HOME_ADDRESS_REVIEW, $0 until HR approval.
  * Multi-store legs -> manual itemisation, $0, route order never guessed.
  * A location NOT on the Matrix -> review, never auto-paid.

Safety: backs the workbook up before writing, saves via a temp file then
atomically swaps it in, and archives a timestamped copy per month. The master
workbook is opened read-only and never modified. If the console is open in
Excel (file locked) the run aborts with a clear message.

Usage:  python process_console.py
"""

import datetime
import os
import shutil
import sys
import tempfile

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation, DataValidationList

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
if SCRIPT_DIR not in sys.path:
    sys.path.insert(0, SCRIPT_DIR)
import console_layout as L
import process_month as pm

WORKFLOW_DIR = os.path.dirname(SCRIPT_DIR)
CONSOLE_PATH = os.path.join(WORKFLOW_DIR, L.CONSOLE_FILENAME)
BACKUP_DIR = os.path.join(WORKFLOW_DIR, "backups")
OUTPUT_DIR = os.path.join(WORKFLOW_DIR, "output")

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

# ---- styling ----
BRAND = "1F4E78"
HDR_FILL = PatternFill("solid", fgColor=BRAND)
HDR_FONT = Font(color="FFFFFF", bold=True)
BORDER = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
KPI_FILL = PatternFill("solid", fgColor="DDEBF7")


# ------------------------------------------------------------------ readers
def _num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        return None


def read_settings(ws):
    """Settings tab (key | value | desc) -> policy dict with sane defaults."""
    policy = {
        "default_rate_per_mile": 0.27, "variance_threshold_pct": 20.0,
        "fuzzy_match_cutoff": 0.72, "fuzzy_match_margin": 0.06, "pay_basis": "matrix",
        "master_workbook": "Fleet_Distance_Dashboard_v5.xlsx",
        "master_matrix_sheet": "Matrix", "master_directory_sheet": "Directory",
        "home_leg_policy": "flag", "multistop_policy": "flag",
    }
    numeric = {"default_rate_per_mile", "variance_threshold_pct",
               "fuzzy_match_cutoff", "fuzzy_match_margin"}
    for r in ws.iter_rows(min_row=2, values_only=True):
        key = (str(r[0]).strip() if r[0] is not None else "")
        if not key:
            continue
        val = r[1] if len(r) > 1 else None
        if key in numeric:
            n = _num(val)
            if n is not None:
                policy[key] = n
        elif val is not None:
            policy[key] = str(val).strip()
    return policy


def read_employees(ws):
    emps = {}
    for r in ws.iter_rows(min_row=2, min_col=1, max_col=5, values_only=True):
        eid = (str(r[0]).strip().replace(".0", "") if r[0] is not None else "")
        if not eid:
            continue
        emps[eid] = {
            "name": ("" if r[1] is None else str(r[1]).strip()),
            "rate": _num(r[2]),
            "active": (str(r[3]).strip().upper() != "NO") if r[3] is not None else True,
        }
    return emps


def read_locations(ws):
    locs = []
    for r in ws.iter_rows(min_row=2, min_col=1, max_col=10, values_only=True):
        loc_id = (str(r[0]).strip() if r[0] is not None else "")
        if not loc_id or loc_id.lower().startswith("(example"):
            continue
        loc = pm.HRLocation()
        loc.loc_id = loc_id
        loc.emp_id = (str(r[1]).strip().replace(".0", "") if r[1] is not None else "") or "ALL"
        loc.type = (str(r[2]).strip().upper() if r[2] is not None else "OTHER") or "OTHER"
        loc.aliases = [pm.normalize(a) for a in str(r[3] or "").split(";") if pm.normalize(a)]
        loc.address = ("" if r[4] is None else str(r[4]).strip())
        loc.lat = _num(r[5]); loc.lon = _num(r[6])
        loc.matrix_code = ("" if r[7] is None else str(r[7]).strip())
        loc.approved = (str(r[8]).strip().upper() == "YES") if r[8] is not None else False
        loc.notes = ("" if r[9] is None else str(r[9]).strip())
        locs.append(loc)
    return locs


def _date_str(v):
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    return str(v).split(" ")[0] if v is not None else ""


def read_monthly_input(ws):
    """Return list of raw leg dicts across ALL months."""
    legs = []
    for r in ws.iter_rows(min_row=2, min_col=1, max_col=9, values_only=True):
        month = (str(r[0]).strip() if r[0] is not None else "")
        date = r[2]
        if not month and (date is None or str(date).strip() == ""):
            continue  # blank row
        legs.append({
            "month": month, "emp_id": (str(r[1]).strip().replace(".0", "") if r[1] is not None else ""),
            "date": date, "origin": str(r[3] or "").strip(), "dest": str(r[4] or "").strip(),
            "notes": str(r[5] or "").strip(), "odos": _num(r[6]), "odoe": _num(r[7]),
            "claimed_miles": _num(r[8]),
        })
    return legs


# ------------------------------------------------------------------ processing
class _Emp:
    def __init__(self, emp_id, name):
        self.emp_id = emp_id
        self.emp_name = name


def process_all(legs, employees, master, policy, hr_locations):
    """Process every leg (all months). Returns list of ResultRow with .month set."""
    default_rate = float(policy.get("default_rate_per_mile", 0.27))
    results = []
    for lg in legs:
        emp = employees.get(lg["emp_id"], {})
        if emp.get("active") is False:
            continue
        rate = emp.get("rate") or default_rate
        shim = _Emp(lg["emp_id"], emp.get("name", ""))
        t = pm.Trip()
        t.date = lg["date"]
        t.origin_raw = lg["origin"]; t.dest_raw = lg["dest"]; t.notes = lg["notes"]
        t.odo_start = lg["odos"]; t.odo_end = lg["odoe"]
        cm = lg["claimed_miles"]
        if cm is None and lg["odos"] is not None and lg["odoe"] is not None:
            cm = lg["odoe"] - lg["odos"]
        t.claimed_miles = cm
        t.claimed_reimb = round(cm * rate, 2) if cm is not None else None
        rr = pm.process_trip(shim, t, master, {}, policy, rate, hr_locations)
        rr.month = lg["month"]
        results.append(rr)
    return results


def _is_review(s):
    return s in pm.REVIEW_STATUSES
def _is_exc(s):
    return s in pm.EXCEPTION_STATUSES


# ------------------------------------------------------------------ writers
def _hdr(ws, headers, row=1):
    for ci, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=ci, value=h)
        c.fill = HDR_FILL; c.font = HDR_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = BORDER
    ws.freeze_panes = f"A{row + 1}"


def _reset_sheet(wb, name):
    """Recreate a sheet in its original position (so tab order is preserved)."""
    if name in wb.sheetnames:
        idx = wb.sheetnames.index(name)
        del wb[name]
    else:
        idx = len(wb.sheetnames)
    return wb.create_sheet(name, idx)


def _autosize(ws, max_w=46):
    for col in ws.columns:
        letter, length = None, 0
        for c in col:
            if letter is None:
                letter = c.column_letter
            if c.value is not None:
                length = max(length, len(str(c.value)))
        if letter:
            ws.column_dimensions[letter].width = min(max(10, length + 2), max_w)


def write_trip_like(ws, rows):
    _hdr(ws, L.TRIP_HEADERS)
    st_col = L.TRIP_HEADERS.index("Status") + 1
    for x in rows:
        supported = "YES" if x.matrix_miles is not None else "NO"
        ws.append([x.month, x.emp_name, x.emp_id, x.date, x.origin_raw, x.dest_raw, x.notes,
                   x.origin_code, x.dest_code, supported, x.matrix_miles, x.ref_miles,
                   x.claimed_miles, x.variance_pct, x.rate, x.reimbursable, x.claimed_reimb,
                   x.status, x.origin_addr, x.dest_addr, x.reason])
        rr = ws.max_row
        fill = pm.STATUS_FILLS.get(x.status)
        if fill:
            ws.cell(row=rr, column=st_col).fill = fill
        for cc in range(1, len(L.TRIP_HEADERS) + 1):
            ws.cell(row=rr, column=cc).border = BORDER
    if not rows:
        ws.cell(row=2, column=1, value="(no trips for the selected month)")
    _autosize(ws)


def write_monthly_report(ws, rows, employees):
    _hdr(ws, L.REPORT_HEADERS)
    ids = []
    for x in rows:
        if x.emp_id not in ids:
            ids.append(x.emp_id)
    grand = [0, 0, 0, 0, 0.0, 0.0, 0.0, 0.0]
    for eid in ids:
        er = [x for x in rows if x.emp_id == eid]
        legs = len(er)
        appr = sum(1 for x in er if x.status == pm.ST_APPROVED)
        rev = sum(1 for x in er if _is_review(x.status))
        exc = sum(1 for x in er if _is_exc(x.status))
        mm = sum(x.matrix_miles for x in er if x.status == pm.ST_APPROVED and x.matrix_miles)
        auto = sum(x.reimbursable for x in er if x.status == pm.ST_APPROVED and x.reimbursable)
        claimed = sum(x.claimed_reimb for x in er if x.claimed_reimb)
        pending = sum((x.claimed_reimb or 0) for x in er if x.status != pm.ST_APPROVED)
        name = er[0].emp_name or employees.get(eid, {}).get("name", "")
        ws.append([name or "(unknown)", eid, legs, appr, rev, exc,
                   round(mm, 1), round(auto, 2), round(claimed, 2), round(pending, 2)])
        for cc in range(1, len(L.REPORT_HEADERS) + 1):
            ws.cell(row=ws.max_row, column=cc).border = BORDER
        for i, v in enumerate([legs, appr, rev, exc, mm, auto, claimed, pending]):
            grand[i] += v
    tr = ws.max_row + 1
    ws.cell(row=tr, column=1, value="TOTAL").font = Font(bold=True)
    for i, v in enumerate([grand[0], grand[1], grand[2], grand[3],
                           round(grand[4], 1), round(grand[5], 2),
                           round(grand[6], 2), round(grand[7], 2)], start=3):
        c = ws.cell(row=tr, column=i, value=v); c.font = Font(bold=True); c.border = BORDER
    if not rows:
        ws.cell(row=2, column=1, value="(no data for the selected month)")
    _autosize(ws)


def write_audit(ws, run_ts, sel_month, policy, employees, hr_locations,
                master, all_results, months):
    _hdr(ws, L.AUDIT_HEADERS)
    approved_locs = sum(1 for l in hr_locations if l.approved)
    rows = [
        ("Run timestamp", run_ts),
        ("Selected month", sel_month),
        ("Config source", L.CONSOLE_FILENAME + " (single HR workbook)"),
        ("Master workbook (read-only)", policy["master_workbook"]),
        ("Pay basis", policy.get("pay_basis")),
        ("Variance threshold %", policy.get("variance_threshold_pct")),
        ("Fuzzy cutoff / margin", f"{policy.get('fuzzy_match_cutoff')} / {policy.get('fuzzy_match_margin')}"),
        ("Employees loaded", len(employees)),
        ("HR verified locations", f"{len(hr_locations)} ({approved_locs} approved)"),
        ("Directory sites / with coords", f"{len(master.codes)} / {len(master.code_to_coord)}"),
        ("Matrix pairs", len(master.matrix)),
        ("Months on file", ", ".join(months)),
        ("Total legs (all months)", len(all_results)),
    ]
    for loc in hr_locations:
        rows.append((f"  HR loc {loc.loc_id}",
                     f"emp {loc.emp_id}, {loc.type}, approved={'YES' if loc.approved else 'NO'}, "
                     f"addr={'set' if loc.address else 'blank'}, coords={'set' if loc.coord else 'blank'}, "
                     f"matrix_code={loc.matrix_code or '-'}"))
    for item, val in rows:
        ws.append([item, val])
        for cc in (1, 2):
            ws.cell(row=ws.max_row, column=cc).border = BORDER
    _autosize(ws, max_w=95)


def update_dashboard(ws, sel_month, run_ts, all_results, months):
    # KPI values for the selected month
    sel = [x for x in all_results if x.month == sel_month]
    kpi = {
        "Total legs": len(sel),
        "Employees": len({x.emp_id for x in sel}),
        "Auto-approved": sum(1 for x in sel if x.status == pm.ST_APPROVED),
        "Needs review": sum(1 for x in sel if _is_review(x.status)),
        "Exceptions": sum(1 for x in sel if _is_exc(x.status)),
        "Auto-approved $": round(sum(x.reimbursable for x in sel
                                     if x.status == pm.ST_APPROVED and x.reimbursable), 2),
        "Pending $": round(sum((x.claimed_reimb or 0) for x in sel
                                if x.status != pm.ST_APPROVED), 2),
    }
    ws[L.DASH_LAST_RUN_CELL] = run_ts
    for i, key in enumerate(L.DASH_KPI_KEYS):
        r = L.DASH_KPI_START_ROW + i
        c = ws.cell(row=r, column=3, value=kpi.get(key))
        c.fill = KPI_FILL; c.border = BORDER

    # all-month history table (clear region then rewrite)
    for rr in range(L.DASH_HISTORY_DATA_ROW,
                    L.DASH_HISTORY_DATA_ROW + L.DASH_HISTORY_MAX_ROWS):
        for cc in range(2, 2 + len(L.DASH_HISTORY_HEADERS)):
            ws.cell(row=rr, column=cc).value = None
    for j, m in enumerate(months):
        mr = [x for x in all_results if x.month == m]
        vals = [m, len(mr),
                sum(1 for x in mr if x.status == pm.ST_APPROVED),
                sum(1 for x in mr if _is_review(x.status)),
                sum(1 for x in mr if _is_exc(x.status)),
                round(sum(x.reimbursable for x in mr if x.status == pm.ST_APPROVED and x.reimbursable), 2),
                round(sum((x.claimed_reimb or 0) for x in mr if x.status != pm.ST_APPROVED), 2)]
        for k, v in enumerate(vals):
            c = ws.cell(row=L.DASH_HISTORY_DATA_ROW + j, column=2 + k, value=v)
            c.border = BORDER

    # refresh the month-selector dropdown to the months actually on file
    ws.data_validations = DataValidationList()
    if months:
        dv = DataValidation(type="list", formula1='"' + ",".join(months) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(L.DASH_SELECTED_MONTH_CELL)


# ------------------------------------------------------------------ main
def main():
    if not os.path.exists(CONSOLE_PATH):
        sys.exit(f"Console not found: {CONSOLE_PATH}\nRun build_console.py first.")

    # Load the console (read-write). If it is open in Excel this still reads,
    # but the later atomic replace will fail -> handled below.
    wb = openpyxl.load_workbook(CONSOLE_PATH)
    for s in (L.SH_SETTINGS, L.SH_EMPLOYEES, L.SH_LOCATIONS, L.SH_INPUT, L.SH_DASHBOARD):
        if s not in wb.sheetnames:
            sys.exit(f"Console is missing the '{s}' tab. Regenerate with build_console.py.")

    policy = read_settings(wb[L.SH_SETTINGS])
    employees = read_employees(wb[L.SH_EMPLOYEES])
    hr_locations = read_locations(wb[L.SH_LOCATIONS])
    legs = read_monthly_input(wb[L.SH_INPUT])

    months = sorted({lg["month"] for lg in legs if lg["month"]})
    if not months:
        sys.exit("No rows found on the Monthly Input tab.")
    sel = wb[L.SH_DASHBOARD][L.DASH_SELECTED_MONTH_CELL].value
    sel_month = str(sel).strip() if sel else ""
    if sel_month not in months:
        note = f"(selected '{sel_month}' not on file; used latest)" if sel_month else "(none selected; used latest)"
        sel_month = months[-1]
        print(f"[info] Month selector {note} -> {sel_month}")

    print(f"[info] Loading master {policy['master_workbook']} (read-only) ...")
    master = pm.MasterData(policy)
    print(f"[info] {len(master.codes)} sites, {len(master.matrix)} matrix pairs, "
          f"{len(master.code_to_coord)} with coords.")
    print(f"[info] Employees={len(employees)}  HR locations={len(hr_locations)}  "
          f"months={months}  selected={sel_month}")

    all_results = process_all(legs, employees, master, policy, hr_locations)
    sel_rows = [x for x in all_results if x.month == sel_month]

    run_ts = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    # ---- write output tabs (recreate) + dashboard (in place) ----
    write_monthly_report(_reset_sheet(wb, L.SH_REPORT), sel_rows, employees)
    write_trip_like(_reset_sheet(wb, L.SH_TRIPS), sel_rows)
    write_trip_like(_reset_sheet(wb, L.SH_EXCEPTIONS),
                    [x for x in sel_rows if x.status != pm.ST_APPROVED])
    write_audit(_reset_sheet(wb, L.SH_AUDIT), run_ts, sel_month, policy,
                employees, hr_locations, master, all_results, months)
    update_dashboard(wb[L.SH_DASHBOARD], sel_month, run_ts, all_results, months)
    # keep tab order stable
    wb._sheets.sort(key=lambda s: L.SHEET_ORDER.index(s.title) if s.title in L.SHEET_ORDER else 99)
    wb.active = wb.sheetnames.index(L.SH_DASHBOARD)

    # ---- safe save: backup -> temp -> atomic replace -> archive ----
    os.makedirs(BACKUP_DIR, exist_ok=True)
    ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = os.path.join(BACKUP_DIR, f"Console_backup_{ts}.xlsx")
    try:
        shutil.copy2(CONSOLE_PATH, backup)
    except Exception as e:
        print(f"[warn] could not write backup: {e}")

    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=WORKFLOW_DIR)
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, CONSOLE_PATH)   # atomic on same volume
    except PermissionError:
        os.path.exists(tmp) and os.remove(tmp)
        sys.exit("ERROR: could not write the console — is it open in Excel? "
                 "Close the workbook and run again.")
    except Exception:
        os.path.exists(tmp) and os.remove(tmp)
        raise

    month_dir = os.path.join(OUTPUT_DIR, sel_month)
    os.makedirs(month_dir, exist_ok=True)
    archive = os.path.join(month_dir, f"Console_{sel_month}_{ts}.xlsx")
    try:
        shutil.copy2(CONSOLE_PATH, archive)
    except Exception as e:
        print(f"[warn] could not archive: {e}")

    # ---- console summary ----
    n_appr = sum(1 for x in sel_rows if x.status == pm.ST_APPROVED)
    n_rev = sum(1 for x in sel_rows if _is_review(x.status))
    n_exc = sum(1 for x in sel_rows if _is_exc(x.status))
    print("\n================ RUN SUMMARY ================")
    print(f"Selected month : {sel_month}")
    print(f"Legs (month)   : {len(sel_rows)}   approved={n_appr} review={n_rev} exceptions={n_exc}")
    print(f"All months     : {', '.join(months)}  (total legs {len(all_results)})")
    print(f"Backup         : {os.path.relpath(backup, WORKFLOW_DIR)}")
    print(f"Archive        : {os.path.relpath(archive, WORKFLOW_DIR)}")
    print(f"Console updated: {CONSOLE_PATH}")
    print("=============================================")


if __name__ == "__main__":
    main()
